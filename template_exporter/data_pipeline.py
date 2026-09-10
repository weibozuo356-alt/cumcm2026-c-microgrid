from __future__ import annotations

import argparse
import gzip
import json
import math
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


SEED = 20260910
START_EXPORT = pd.Timestamp("2025-02-01")
END_EXPORT = pd.Timestamp("2025-12-31")
N_INTERVALS = 144
N_SCENARIOS = 20


def fmt_source_label(value: object) -> str:
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    text = str(value).strip()
    if text in {"0:00+1", "00:00+1", "0:00:00+1", "00:00:00+1"}:
        return "00:00+1"
    try:
        return pd.to_datetime(text).strftime("%H:%M:%S")
    except Exception:
        return text


def season(month: int) -> str:
    if month in (12, 1, 2):
        return "winter"
    if month in (3, 4, 5):
        return "spring"
    if month in (6, 7, 8):
        return "summer"
    return "autumn"


def logical_target(ts: pd.Timestamp) -> tuple[pd.Timestamp, int]:
    minute = ts.hour * 60 + ts.minute
    if minute == 0:
        return ts.normalize() - pd.Timedelta(days=1), 144
    return ts.normalize(), minute // 10


def read_output_labels(template: Path) -> list[str]:
    wb = load_workbook(template, read_only=True, data_only=False)
    ws = wb["计划购电量"]
    labels = [str(ws.cell(1, c).value) for c in range(2, 146)]
    if len(labels) != N_INTERVALS:
        raise ValueError(f"模板区间列数不是144: {template}")
    return labels


def wide_sheet(path: Path, sheet: str | int = 0) -> tuple[pd.DataFrame, list[str], list[object]]:
    raw = pd.read_excel(path, sheet_name=sheet, dtype=object)
    raw = raw.rename(columns={raw.columns[0]: "date"})
    raw["date"] = pd.to_datetime(raw["date"], errors="coerce").dt.normalize()
    time_cols = list(raw.columns[1:])
    vals = raw[time_cols].apply(pd.to_numeric, errors="coerce")
    vals.index = raw["date"]
    vals.columns = range(1, len(time_cols) + 1)
    return vals, [fmt_source_label(x) for x in time_cols], time_cols


def attachment3_long(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    raw = pd.read_excel(path, dtype=object)
    raw.columns = [str(x).strip() for x in raw.columns]
    raw["_date_was_blank"] = raw["日期"].isna() | (raw["日期"].astype(str).str.strip() == "")
    raw["日期"] = raw["日期"].replace("", np.nan).ffill()
    raw["issue_date"] = pd.to_datetime(raw["日期"], errors="coerce").dt.normalize()
    raw["issue_hour"] = pd.to_numeric(raw["预报时刻"].astype(str).str.extract(r"(\d+)")[0], errors="coerce")
    lead_cols = [c for c in raw.columns if str(c).startswith("预报") and c != "预报时刻"]
    records: list[dict[str, object]] = []
    for _, row in raw.iterrows():
        if pd.isna(row["issue_date"]) or pd.isna(row["issue_hour"]):
            continue
        issue_ts = row["issue_date"] + pd.Timedelta(hours=int(row["issue_hour"]))
        for lead, col in enumerate(lead_cols, 1):
            value = pd.to_numeric(pd.Series([row[col]]), errors="coerce").iloc[0]
            for substep in range(1, 7):
                target_ts = issue_ts + pd.Timedelta(hours=lead - 1, minutes=10 * substep)
                target_date, interval_index = logical_target(target_ts)
                records.append(
                    {
                        "forecast_issue_time": issue_ts,
                        "issue_date": row["issue_date"],
                        "issue_hour": int(row["issue_hour"]),
                        "forecast_lead_h": lead,
                        "substep_10min": substep,
                        "target_timestamp": target_ts,
                        "date": target_date,
                        "interval_index": interval_index,
                        "pv_forecast_kw": float(value) if pd.notna(value) else np.nan,
                    }
                )
    mapped = pd.DataFrame(records)
    return raw, mapped


def expanding_mean(a: np.ndarray) -> np.ndarray:
    out = np.full_like(a, np.nan, dtype=float)
    sums = np.zeros(a.shape[1], dtype=float)
    counts = np.zeros(a.shape[1], dtype=float)
    for i in range(a.shape[0]):
        out[i] = np.divide(sums, counts, out=np.full(a.shape[1], np.nan), where=counts > 0)
        ok = np.isfinite(a[i])
        sums[ok] += a[i, ok]
        counts[ok] += 1
    return out


def grouped_expanding_mean(a: np.ndarray, groups: np.ndarray, fallback: np.ndarray) -> np.ndarray:
    out = np.full_like(a, np.nan, dtype=float)
    sums = {0: np.zeros(a.shape[1]), 1: np.zeros(a.shape[1])}
    counts = {0: np.zeros(a.shape[1]), 1: np.zeros(a.shape[1])}
    for i in range(a.shape[0]):
        g = int(groups[i])
        pred = np.divide(sums[g], counts[g], out=np.full(a.shape[1], np.nan), where=counts[g] > 0)
        out[i] = np.where(np.isfinite(pred), pred, fallback[i])
        ok = np.isfinite(a[i])
        sums[g][ok] += a[i, ok]
        counts[g][ok] += 1
    return out


def ridge_expanding(a: np.ndarray, dates: pd.DatetimeIndex, fallback: np.ndarray, nonnegative: bool) -> np.ndarray:
    out = np.full_like(a, np.nan, dtype=float)
    p = 8
    xtx = np.zeros((p, p))
    xty = np.zeros(p)
    scale = 10000.0
    interval = np.arange(1, N_INTERVALS + 1)
    time_sin = np.sin(2 * np.pi * interval / N_INTERVALS)
    time_cos = np.cos(2 * np.pi * interval / N_INTERVALS)

    for i, d in enumerate(dates):
        if i < 7:
            out[i] = fallback[i]
            continue
        doy = d.dayofyear
        x = np.column_stack(
            [
                a[i - 1] / scale,
                a[i - 7] / scale,
                time_sin,
                time_cos,
                np.full(N_INTERVALS, np.sin(2 * np.pi * doy / 365.0)),
                np.full(N_INTERVALS, np.cos(2 * np.pi * doy / 365.0)),
                np.full(N_INTERVALS, float(d.dayofweek >= 5)),
                np.ones(N_INTERVALS),
            ]
        )
        if i > 7:
            reg = np.eye(p) * 1e-2
            reg[-1, -1] = 1e-8
            try:
                beta = np.linalg.solve(xtx + reg, xty)
                pred = x @ beta * scale
            except np.linalg.LinAlgError:
                pred = fallback[i]
            out[i] = pred
        else:
            out[i] = fallback[i]

        y = a[i] / scale
        ok = np.isfinite(x).all(axis=1) & np.isfinite(y)
        xtx += x[ok].T @ x[ok]
        xty += x[ok].T @ y[ok]

    if nonnegative:
        out = np.maximum(out, 0.0)
    return out


def bias_correct(raw: np.ndarray, actual: np.ndarray) -> np.ndarray:
    out = np.full_like(raw, np.nan, dtype=float)
    sums = np.zeros(raw.shape[1])
    counts = np.zeros(raw.shape[1])
    for i in range(raw.shape[0]):
        bias = np.divide(sums, counts, out=np.zeros(raw.shape[1]), where=counts > 0)
        out[i] = np.maximum(raw[i] + bias, 0.0)
        resid = actual[i] - raw[i]
        ok = np.isfinite(resid)
        sums[ok] += resid[ok]
        counts[ok] += 1
    return out


def metric_rows(actual: np.ndarray, pred: np.ndarray, dates: np.ndarray, interval: np.ndarray,
                variable: str, model: str) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []

    def add(kind: str, value: str, mask: np.ndarray) -> None:
        ok = mask & np.isfinite(actual) & np.isfinite(pred)
        if not ok.any():
            return
        err = pred[ok] - actual[ok]
        denom = np.abs(actual[ok]).sum()
        rows.append(
            {
                "variable": variable,
                "model": model,
                "segment_type": kind,
                "segment_value": value,
                "n": int(ok.sum()),
                "mae": float(np.mean(np.abs(err))),
                "rmse": float(np.sqrt(np.mean(err ** 2))),
                "wmape": float(np.abs(err).sum() / denom) if denom > 0 else np.nan,
            }
        )

    base = (dates >= np.datetime64(START_EXPORT)) & (dates <= np.datetime64(END_EXPORT))
    add("overall", "all", base)
    months = pd.DatetimeIndex(dates).month.to_numpy()
    for m in range(2, 13):
        add("month", f"{m:02d}", base & (months == m))
    if variable == "pv":
        daylight = actual > 1.0
        add("daylight", "actual_pv_gt_1kw", base & daylight)
    lead = np.ceil(interval / 6).astype(int)
    for h in range(1, 25):
        add("lead_h", str(h), base & (lead == h))
    return rows


def make_forecasts(processed: pd.DataFrame, dates: pd.DatetimeIndex, load: pd.DataFrame,
                   pv: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    load_a = load.to_numpy(float)
    pv_a = pv.to_numpy(float)
    exp_load = expanding_mean(load_a)
    exp_pv = expanding_mean(pv_a)
    weekend = (dates.dayofweek >= 5).astype(int)

    preds: dict[str, np.ndarray] = {}
    for name, arr, exp, nonnegative in (
        ("load", load_a, exp_load, False),
        ("pv", pv_a, exp_pv, True),
    ):
        prev = np.vstack([np.full((1, N_INTERVALS), np.nan), arr[:-1]])
        week = np.vstack([np.full((7, N_INTERVALS), np.nan), arr[:-7]])
        grouped = grouped_expanding_mean(arr, weekend, exp)
        ridge = ridge_expanding(arr, dates, exp, nonnegative)
        preds[f"{name}_pred_prev_day_kw"] = np.maximum(prev, 0) if nonnegative else prev
        preds[f"{name}_pred_prev_week_kw"] = np.maximum(week, 0) if nonnegative else week
        preds[f"{name}_pred_expanding_mean_kw"] = np.maximum(exp, 0) if nonnegative else exp
        preds[f"{name}_pred_daytype_mean_kw"] = np.maximum(grouped, 0) if nonnegative else grouped
        preds[f"{name}_pred_ridge_kw"] = ridge

    raw_official = processed.pivot(index="date", columns="interval_index", values="pv_forecast_kw").reindex(dates).to_numpy(float)
    preds["pv_pred_official_kw"] = raw_official
    preds["pv_pred_official_bias_corrected_kw"] = bias_correct(raw_official, pv_a)

    def causal_select(actual: np.ndarray, candidates: list[str], default: str) -> tuple[np.ndarray, np.ndarray]:
        selected = np.full_like(actual, np.nan, dtype=float)
        chosen = np.empty(actual.shape, dtype=object)
        err_sum = {c: 0.0 for c in candidates}
        err_count = {c: 0 for c in candidates}
        for i in range(actual.shape[0]):
            eligible = [c for c in candidates if err_count[c] > 0 and np.isfinite(preds[c][i]).all()]
            pick = min(eligible, key=lambda c: err_sum[c] / err_count[c]) if eligible else default
            selected[i] = preds[pick][i]
            chosen[i, :] = pick
            for c in candidates:
                ok = np.isfinite(actual[i]) & np.isfinite(preds[c][i])
                if ok.any():
                    err_sum[c] += float(np.abs(actual[i, ok] - preds[c][i, ok]).sum())
                    err_count[c] += int(ok.sum())
        return selected, chosen

    load_candidates = [
        "load_pred_prev_day_kw", "load_pred_prev_week_kw", "load_pred_expanding_mean_kw",
        "load_pred_daytype_mean_kw", "load_pred_ridge_kw",
    ]
    pv_candidates = [
        "pv_pred_prev_day_kw", "pv_pred_prev_week_kw", "pv_pred_expanding_mean_kw",
        "pv_pred_daytype_mean_kw", "pv_pred_ridge_kw", "pv_pred_official_kw",
        "pv_pred_official_bias_corrected_kw",
    ]
    load_selected, load_chosen = causal_select(load_a, load_candidates, "load_pred_expanding_mean_kw")
    pv_selected, pv_chosen = causal_select(pv_a, pv_candidates, "pv_pred_official_kw")

    base = processed[["date", "interval_index", "source_time_label", "output_interval_label"]].copy()
    base = base.sort_values(["date", "interval_index"]).reset_index(drop=True)
    base["load_actual_kw"] = load_a.reshape(-1)
    base["pv_actual_kw"] = pv_a.reshape(-1)
    for col, arr in preds.items():
        base[col] = arr.reshape(-1)
    base["load_pred_selected_kw"] = load_selected.reshape(-1)
    base["load_selected_model"] = load_chosen.reshape(-1)
    base["pv_pred_selected_kw"] = pv_selected.reshape(-1)
    base["pv_selected_model"] = pv_chosen.reshape(-1)
    base["load_residual_kw"] = base["load_actual_kw"] - base["load_pred_selected_kw"]
    base["pv_residual_kw"] = base["pv_actual_kw"] - base["pv_pred_selected_kw"]
    base["net_load_actual_kw"] = base["load_actual_kw"] - base["pv_actual_kw"]
    base["net_load_pred_kw"] = base["load_pred_selected_kw"] - base["pv_pred_selected_kw"]
    base["net_load_residual_kw"] = base["net_load_actual_kw"] - base["net_load_pred_kw"]
    base["forecast_available_at"] = base["date"].dt.strftime("%Y-%m-%d") + " 00:00:00"

    metrics: list[dict[str, object]] = []
    dflat = np.repeat(dates.to_numpy(dtype="datetime64[ns]"), N_INTERVALS)
    iflat = np.tile(np.arange(1, N_INTERVALS + 1), len(dates))
    for col in [c for c in base.columns if c.startswith("load_pred_") and c.endswith("_kw") and "selected" not in c]:
        metrics.extend(metric_rows(base["load_actual_kw"].to_numpy(), base[col].to_numpy(), dflat, iflat, "load", col))
    for col in [c for c in base.columns if c.startswith("pv_pred_") and c.endswith("_kw") and "selected" not in c]:
        metrics.extend(metric_rows(base["pv_actual_kw"].to_numpy(), base[col].to_numpy(), dflat, iflat, "pv", col))
    metrics.extend(metric_rows(base["load_actual_kw"].to_numpy(), base["load_pred_selected_kw"].to_numpy(), dflat, iflat, "load", "load_pred_causal_selected_kw"))
    metrics.extend(metric_rows(base["pv_actual_kw"].to_numpy(), base["pv_pred_selected_kw"].to_numpy(), dflat, iflat, "pv", "pv_pred_causal_selected_kw"))
    return base, pd.DataFrame(metrics)


def official_issue_metrics(issue_long: pd.DataFrame, processed: pd.DataFrame) -> pd.DataFrame:
    actual = processed[["date", "interval_index", "pv_actual_kw"]]
    x = issue_long.merge(actual, on=["date", "interval_index"], how="inner")
    x = x[(x["date"] >= START_EXPORT) & (x["date"] <= END_EXPORT)].copy()
    rows: list[dict[str, object]] = []

    def add(segment_type: str, segment_value: str, g: pd.DataFrame) -> None:
        ok = g["pv_actual_kw"].notna() & g["pv_forecast_kw"].notna()
        y = g.loc[ok, "pv_actual_kw"].to_numpy(float)
        p = g.loc[ok, "pv_forecast_kw"].to_numpy(float)
        if not len(y):
            return
        err = p - y
        denom = np.abs(y).sum()
        rows.append({
            "variable": "pv",
            "model": "pv_pred_official_all_issues_kw",
            "segment_type": segment_type,
            "segment_value": segment_value,
            "n": int(len(y)),
            "mae": float(np.mean(np.abs(err))),
            "rmse": float(np.sqrt(np.mean(err ** 2))),
            "wmape": float(np.abs(err).sum() / denom) if denom > 0 else np.nan,
        })

    add("overall_all_issues", "all", x)
    for issue, g in x.groupby("issue_hour"):
        add("issue_hour", str(int(issue)), g)
    for (issue, lead), g in x.groupby(["issue_hour", "forecast_lead_h"]):
        add("issue_hour_lead", f"H{int(issue):02d}_L{int(lead):02d}", g)
    return pd.DataFrame(rows)


def write_scenarios(root: Path, forecasts: pd.DataFrame, processed: pd.DataFrame) -> None:
    out = root / "scenario_data"
    out.mkdir(exist_ok=True)
    residuals = forecasts.merge(processed[["date", "interval_index", "price_actual"]], on=["date", "interval_index"], how="left")
    residuals = residuals.sort_values(["date", "interval_index"]).reset_index(drop=True)
    residuals["price_pred_prev_day"] = residuals.groupby("interval_index")["price_actual"].shift(1)
    residuals["price_residual"] = residuals["price_actual"] - residuals["price_pred_prev_day"]
    keep = [
        "date", "interval_index", "output_interval_label", "load_actual_kw", "load_pred_selected_kw",
        "load_residual_kw", "pv_actual_kw", "pv_pred_selected_kw", "pv_residual_kw",
        "net_load_actual_kw", "net_load_pred_kw", "net_load_residual_kw", "price_actual",
        "price_pred_prev_day", "price_residual",
    ]
    residuals[keep].to_csv(out / "residuals_long.csv", index=False, encoding="utf-8-sig", float_format="%.6f")

    traj = residuals.pivot(index="date", columns="interval_index", values=["load_residual_kw", "pv_residual_kw", "net_load_residual_kw", "price_residual"])
    traj.columns = [f"{a}_{int(b):03d}" for a, b in traj.columns]
    traj.reset_index().to_csv(out / "daily_residual_trajectories.csv.gz", index=False, compression="gzip", encoding="utf-8", float_format="%.6f")

    required_residuals = ["load_residual_kw", "pv_residual_kw", "net_load_residual_kw", "price_residual"]
    daily = {
        d: g.sort_values("interval_index")
        for d, g in residuals.groupby("date")
        if len(g) == N_INTERVALS and g[required_residuals].notna().all().all()
    }
    target_days = pd.date_range(START_EXPORT, END_EXPORT, freq="D")
    rng = np.random.default_rng(SEED)
    scenario_path = out / "joint_scenarios.csv.gz"
    with gzip.open(scenario_path, "wt", encoding="utf-8", newline="") as fh:
        header = True
        for target in target_days:
            candidates = [d for d in daily if d < target and season(d.month) == season(target.month) and (d.dayofweek >= 5) == (target.dayofweek >= 5)]
            if not candidates:
                candidates = [d for d in daily if d < target]
            chosen = rng.choice(np.array(candidates, dtype="datetime64[ns]"), size=N_SCENARIOS, replace=True)
            chunks = []
            for sid, src in enumerate(chosen, 1):
                src_day = pd.Timestamp(src)
                g = daily[src_day][["interval_index", "load_residual_kw", "pv_residual_kw", "net_load_residual_kw", "price_residual"]].copy()
                g.insert(0, "source_date", src_day.strftime("%Y-%m-%d"))
                g.insert(0, "scenario_id", sid)
                g.insert(0, "target_date", target.strftime("%Y-%m-%d"))
                chunks.append(g)
            pd.concat(chunks, ignore_index=True).to_csv(fh, index=False, header=header, float_format="%.6f")
            header = False

    sample = pd.read_csv(scenario_path, compression="gzip", nrows=N_SCENARIOS * N_INTERVALS)
    sample.to_csv(out / "joint_scenarios_sample.csv", index=False, encoding="utf-8-sig")

    margin_rows = []
    residual_arrays = {
        d: g.sort_values("interval_index")[["load_residual_kw", "pv_residual_kw", "net_load_residual_kw", "price_residual"]].to_numpy(float)
        for d, g in daily.items()
    }
    for target in target_days:
        candidates = [d for d in residual_arrays if d < target and season(d.month) == season(target.month) and (d.dayofweek >= 5) == (target.dayofweek >= 5)]
        if len(candidates) < 3:
            candidates = [d for d in residual_arrays if d < target]
        cube = np.stack([residual_arrays[d] for d in candidates])
        qs = np.nanquantile(cube, [0.05, 0.50, 0.90, 0.95], axis=0)
        for idx in range(N_INTERVALS):
            row: dict[str, object] = {
                "target_date": target.strftime("%Y-%m-%d"),
                "interval_index": idx + 1,
                "history_days": len(candidates),
                "season": season(target.month),
                "day_type": "weekend" if target.dayofweek >= 5 else "weekday",
            }
            for v, name in enumerate(["load", "pv", "net_load", "price"]):
                for qi, qname in enumerate(["q05", "q50", "q90", "q95"]):
                    row[f"{name}_{qname}"] = qs[qi, idx, v]
            margin_rows.append(row)
    pd.DataFrame(margin_rows).to_csv(out / "residual_quantile_margins.csv.gz", index=False, compression="gzip", encoding="utf-8", float_format="%.6f")


def latest_issue_forecast(issue_long: pd.DataFrame, dates: pd.DatetimeIndex) -> pd.DataFrame:
    x = issue_long[(issue_long["date"].isin(dates)) & (issue_long["issue_date"] == issue_long["date"])].copy()
    x = x.sort_values(["date", "interval_index", "forecast_issue_time"])
    x = x.drop_duplicates(["date", "interval_index"], keep="last")
    return x[["date", "interval_index", "pv_forecast_kw", "forecast_issue_time"]].rename(columns={"pv_forecast_kw": "pv_forecast_latest_kw"})


def make_provisional_solution(root: Path, processed: pd.DataFrame, forecasts: pd.DataFrame,
                              issue_long: pd.DataFrame, labels: list[str]) -> pd.DataFrame:
    data = forecasts.merge(processed[["date", "interval_index", "price_actual", "price_static"]], on=["date", "interval_index"], how="left")
    data = data[(data["date"] >= START_EXPORT) & (data["date"] <= END_EXPORT)].copy()
    latest = latest_issue_forecast(issue_long, pd.date_range(START_EXPORT, END_EXPORT))
    data = data.merge(latest, on=["date", "interval_index"], how="left")
    data["pv_forecast_latest_kw"] = data["pv_forecast_latest_kw"].fillna(data["pv_pred_official_bias_corrected_kw"])

    rows = []
    for strategy, price_col, adjusted in (
        ("result2", "price_static", False),
        ("result3", "price_static", True),
        ("result4-2", "price_actual", False),
        ("result4-3", "price_actual", True),
    ):
        z = data.copy()
        plan = np.maximum(z["load_pred_selected_kw"] - z["pv_pred_official_bias_corrected_kw"], 0.0) / 6.0
        adjusted_qty = np.maximum(z["load_pred_selected_kw"] - z["pv_forecast_latest_kw"], 0.0) / 6.0 if adjusted else plan.copy()
        actual_need = np.maximum(z["load_actual_kw"] - z["pv_actual_kw"], 0.0) / 6.0
        emergency = np.maximum(actual_need - adjusted_qty, 0.0)
        price = z[price_col].to_numpy(float)
        cost_plan = plan.to_numpy(float) * price
        diff = adjusted_qty.to_numpy(float) - plan.to_numpy(float)
        cost_adjust = np.where(diff >= 0, 1.5 * price * diff, 0.5 * price * (-diff)) if adjusted else np.zeros(len(z))
        cost_emergency = 5.0 * price * emergency.to_numpy(float)
        final = adjusted_qty.to_numpy(float) + emergency.to_numpy(float)
        part = pd.DataFrame(
            {
                "date": z["date"].dt.strftime("%Y-%m-%d"),
                "interval_index": z["interval_index"].astype(int),
                "output_interval_label": z["output_interval_label"],
                "purchase_plan_kwh": plan,
                "purchase_adjust_kwh": adjusted_qty,
                "purchase_final_kwh": final,
                "purchase_emergency_kwh": emergency,
                "charge_kwh": 0.0,
                "discharge_kwh": 0.0,
                "soc_start_kwh": 6000.0,
                "soc_end_kwh": 6000.0,
                "curtailment_kwh": np.maximum(z["pv_actual_kw"] - z["load_actual_kw"], 0.0) / 6.0,
                "cost_plan": cost_plan,
                "cost_adjust": cost_adjust,
                "cost_emergency": cost_emergency,
                "cost_total": cost_plan + cost_adjust + cost_emergency,
                "strategy_name": strategy,
                "solver_status": "PROVISIONAL_CAUSAL_NO_STORAGE",
            }
        )
        rows.append(part)

    att1 = pd.read_excel(root / "C题" / "附件" / "附件1.xlsx")
    load1 = pd.to_numeric(att1.iloc[:, 2], errors="coerce").to_numpy(float)
    pv1 = pd.to_numeric(att1.iloc[:, 3], errors="coerce").to_numpy(float)
    price1 = pd.to_numeric(att1.iloc[:, 1], errors="coerce").to_numpy(float)
    plan1 = np.maximum(load1 - pv1, 0.0) / 6.0
    q1 = pd.DataFrame(
        {
            "date": "2025-01-01",
            "interval_index": np.arange(1, N_INTERVALS + 1),
            "output_interval_label": labels,
            "purchase_plan_kwh": plan1,
            "purchase_adjust_kwh": plan1,
            "purchase_final_kwh": plan1,
            "purchase_emergency_kwh": 0.0,
            "charge_kwh": 0.0,
            "discharge_kwh": 0.0,
            "soc_start_kwh": 6000.0,
            "soc_end_kwh": 6000.0,
            "curtailment_kwh": np.maximum(pv1 - load1, 0.0) / 6.0,
            "cost_plan": plan1 * price1,
            "cost_adjust": 0.0,
            "cost_emergency": 0.0,
            "cost_total": plan1 * price1,
            "strategy_name": "result1",
            "solver_status": "PROVISIONAL_Q1_NO_STORAGE",
        }
    )
    rows.insert(0, q1)
    result = pd.concat(rows, ignore_index=True)
    result.to_csv(root / "scenario_data" / "provisional_solution_long.csv", index=False, encoding="utf-8-sig", float_format="%.6f")
    return result


def audit_markdown(root: Path, audit: dict[str, object]) -> None:
    text = f"""# 数据审计报告

生成时间：{datetime.now().isoformat(timespec='seconds')}（Asia/Shanghai）

## 文件结构

| 文件 | 工作表 | 原始形状 | 日期范围 | 结论 |
|---|---|---:|---|---|
| 附件1.xlsx | Sheet1 | 145×4 | 单日典型曲线 | 144 个十分钟记录 |
| 附件2.xlsx | 小区负载、光伏发电实际功率 | 各 366×145 | 2025-01-01—2025-12-31 | 每日 144 个记录 |
| 附件3.xlsx | Sheet1 | 1461×26 | 2025-01-01—2025-12-31 | 每日 4 次发布、每次 24 个整点提前期 |
| 附件4.xlsx | Sheet1 | 366×145 | 2025-01-01—2025-12-31 | 每日 144 个动态电价 |

## 自动审计结果

```json
{json.dumps(audit, ensure_ascii=False, indent=2)}
```

## 时间与语义核对

| 类别 | 事实/解释 |
|---|---|
| 题面事实 | 附件 2/4 的列为每日十分钟数据；附件 3 在每日 0:00、6:00、12:00、18:00 发布未来 24 小时整点光伏预测。附件 1 中 `0:00+1` 表示次日 0:00。 |
| 官方模板事实 | 年度模板首个输出列为 `0:10-0:20`，第144个输出列为 `0:00-0:10+1`；result1.xlsx 的第144行写作 `0:00+1-0:10+1`，两者文字不完全一致。 |
| 采用的解释 | 严格遵从团队口径：附件宽表第 t 个数绑定 `interval_index=t`，并写入模板第 t 个位置，不整体平移。附件 3 的“预报 h 小时”按该小时内 6 个十分钟区间的分段常数功率展开。 |
| 敏感性处理 | `scenario_data/official_forecasts_long.csv.gz` 保留发布时刻、提前期、目标时间和十分钟子步，后续可替换为线性插值并复算；当前主结果不做插值。 |

## 质量规则

- 负荷、光伏、价格全部强制转为数值；原始文本数字数量已计入 JSON。
- 负荷与光伏单位为 kW，静态/动态电价单位为元/kWh；十分钟电量换算系数为 1/6。
- 附件 3 的空日期仅是合并展示语义，已向下填充；预测目标跨日、跨年由真实时间戳计算。
- 附件3类型检查中的1460个非数值文本均为`预报时刻`（0:00/6:00/12:00/18:00），不是脏数值。
- 未发现缺失、重复日期/时标、负负荷、负光伏或负电价。报告保留极值及0.1%/99.9%分位数；没有因全局统计“离群”而删除季节性峰值。
- 预测与指标均按时间顺序生成；没有随机打乱，也没有用目标日未来真实值训练。

## 官方论文与支撑材料要求

- 电子版论文建议 PDF，首页必须为摘要专用页，不含承诺书和编号页，文件不超过 20 MB。
- 正文不含目录且不超过 30 页；附录页数不限，必须列出支撑材料文件清单和全部可运行源代码。
- 支撑材料须打包为 RAR/ZIP，不超过 20 MB；不得含参赛者、学校或赛区身份信息。
- 数据、程序、结果与论文数字必须一致，否则存在取消评奖资格风险。

## 已知风险

- GitHub 团队仓库在 2026-09-10 核对时为空；本次仅能采用用户消息中给出的 Q1→Q2→Q3→Q4 主线。
- “调整购电量”字段的绝对量/增量语义题面未以接口字段明示；本管线采用“调整后的绝对计划购电量”，并在交接文档中显式标注。
- 官方模板的显示区间相对附件时间标签存在 10 分钟表面差异；依据用户明确口径按序号一一对应。
- result1.xlsx模板本身没有“全天购电量/全天购电费”列；为满足“不修改宽表表头”，导出器不新增列，相关合计由验证报告从144个区间独立计算。
"""
    (root / "data_audit.md").write_text(text, encoding="utf-8")


def data_dictionary(root: Path) -> None:
    text = """# 数据字典

## processed_timeseries.csv

| 列 | 类型 | 单位 | 来源 | 最早可获得时刻/说明 |
|---|---|---|---|---|
| date | date | 日 | 附件2 | 逻辑运行日 |
| interval_index | int | 1—144 | 附件列序 | 内部唯一时间索引 |
| source_time_label | string | 时刻 | 附件2表头 | 原始标签；第144项为次日0:00 |
| output_interval_label | string | 区间 | 官方模板 | 第t项与interval_index=t严格对应 |
| interval_end_timestamp | datetime | 时刻 | 计算 | date + 10×interval_index分钟 |
| load_actual_kw | float | kW | 附件2/小区负载 | 区间结束后可获得；不得用于该区间事前决策 |
| pv_actual_kw | float | kW | 附件2/光伏实际 | 区间结束后可获得；不得用于该区间事前决策 |
| price_actual | float | 元/kWh | 附件4 | 动态电价；可获得时刻由Q4建模假设决定 |
| price_static | float | 元/kWh | 附件1 | Q1—Q3每日相同电价曲线 |
| forecast_issue_time | datetime | 时刻 | 附件3 | 当日0:00版本，用于全天基准计划 |
| forecast_lead_h | int | 小时 | 附件3列名 | ceil(interval_index/6) |
| pv_forecast_kw | float | kW | 附件3 | 对应0:00发布版本，分段常数展开到10分钟 |

## forecast_results.csv

包含实际值、前一日同刻、上周同日同刻、历史扩展同刻均值、工作日/周末扩展均值、扩展窗口岭回归、官方光伏预测和官方预测偏差校正。`*_selected_*`每天只根据前一日结束前的累计MAE选择候选模型，`*_selected_model`记录选择结果；它们是场景残差默认基准。所有预测和模型选择只使用目标时刻之前已获得的信息。

## solution_long.csv 接口

| 列 | 单位/含义 |
|---|---|
| date | 逻辑运行日 |
| interval_index | 1—144 |
| output_interval_label | 官方模板显示区间，必须与索引匹配 |
| purchase_plan_kwh | 0:00确定的计划购电量，kWh |
| purchase_adjust_kwh | 采用的接口解释：更新预测后调整得到的绝对计划购电量，非差值，kWh；无调整问题等于purchase_plan_kwh |
| purchase_final_kwh | 最终实际外部购电量，kWh；Q3口径为purchase_adjust_kwh + purchase_emergency_kwh |
| purchase_emergency_kwh | 紧急购电量，kWh |
| charge_kwh / discharge_kwh | 交流侧充/放电电量，kWh |
| soc_start_kwh / soc_end_kwh | 区间起/末储电量，kWh |
| curtailment_kwh | 弃光电量，kWh |
| cost_plan / cost_adjust / cost_emergency / cost_total | 区间费用，元 |
| strategy_name | `result1`、`result2`、`result3`、`result4-2`、`result4-3`之一 |
| solver_status | 求解器状态；最终官方文件不得是PROVISIONAL状态 |

## 储能统一口径

- 容量 12000 kWh；允许储电量 1200—10800 kWh；初值 6000 kWh。
- 最大充/放电功率 5000 kW，对应每十分钟 833.3333 kWh。
- 主模型充、放电效率均为0.9；建议状态式 `soc_end = soc_start + 0.9*charge_kwh - discharge_kwh/0.9`。
- Q2—Q4跨日连续，不得每日重置。
"""
    (root / "data_dictionary.md").write_text(text, encoding="utf-8")


def scenario_description(root: Path) -> None:
    text = f"""# 不确定性场景说明

- 随机种子：`{SEED}`。
- 默认点预测：负荷与光伏均在候选模型间按截至前一日的累计MAE进行因果选择；动态电价采用前一日同刻。
- 残差：实际值减点预测值；净负荷残差严格等于负荷残差减光伏残差。
- `residuals_long.csv`：逐十分钟负荷、光伏、净负荷和动态电价残差。
- `daily_residual_trajectories.csv.gz`：完整日残差轨迹，保留日内相关性。
- `joint_scenarios.csv.gz`：针对每个2025-02-01至2025-12-31目标日，从严格早于目标日的同季节、同工作日类型完整日轨迹中有放回抽样；每个目标日 {N_SCENARIOS} 个联合净负荷—电价场景。历史不足时退化为全部既往日。
- `joint_scenarios_sample.csv`：首个目标日的可读样例。
- `residual_quantile_margins.csv.gz`：对每个目标日和区间，仅用目标日前历史计算5%、50%、90%、95%分位数，可直接构造安全裕度。
- 附件3预测按每个整点提前期覆盖随后6个十分钟子区间的分段常数解释展开；所有发布版本与跨日目标均保存在`official_forecasts_long.csv.gz`。
"""
    (root / "scenario_description.md").write_text(text, encoding="utf-8")


def build_audit(attach: Path, load: pd.DataFrame, pv: pd.DataFrame, price: pd.DataFrame,
                raw3: pd.DataFrame, issue_long: pd.DataFrame, source_labels: list[str]) -> dict[str, object]:
    def stat(df: pd.DataFrame) -> dict[str, object]:
        arr = df.to_numpy(float)
        return {
            "days": int(len(df)),
            "intervals_per_day": int(df.shape[1]),
            "date_min": str(df.index.min().date()),
            "date_max": str(df.index.max().date()),
            "missing": int(np.isnan(arr).sum()),
            "negative": int((arr < 0).sum()),
            "min": float(np.nanmin(arr)),
            "max": float(np.nanmax(arr)),
            "q0_1_percent": float(np.nanquantile(arr, 0.001)),
            "q99_9_percent": float(np.nanquantile(arr, 0.999)),
            "duplicate_dates": int(df.index.duplicated().sum()),
        }

    def cell_types(path: Path, sheet: str | int) -> dict[str, int]:
        raw = pd.read_excel(path, sheet_name=sheet, dtype=object).iloc[:, 1:]
        values = raw.to_numpy(object).reshape(-1)
        text_numeric = 0
        nonnumeric_text = 0
        for value in values:
            if isinstance(value, str) and value.strip():
                try:
                    float(value)
                    text_numeric += 1
                except ValueError:
                    nonnumeric_text += 1
        return {"text_numeric_cells": text_numeric, "nonnumeric_text_cells": nonnumeric_text}

    att1 = pd.read_excel(attach / "附件1.xlsx", dtype=object)
    att1_num = att1.iloc[:, 1:].apply(pd.to_numeric, errors="coerce").to_numpy(float)
    issue_counts = raw3.groupby("issue_date")["issue_hour"].agg(["count", "nunique"])
    lead_cols = [c for c in raw3.columns if str(c).startswith("预报") and c != "预报时刻"]
    lead_values = raw3[lead_cols].apply(pd.to_numeric, errors="coerce").to_numpy(float)
    return {
        "attachment1": {
            "rows": int(len(att1)),
            "columns": int(att1.shape[1]),
            "time_labels_unique": int(att1.iloc[:, 0].astype(str).nunique()),
            "first_time_label": fmt_source_label(att1.iloc[0, 0]),
            "last_time_label": fmt_source_label(att1.iloc[-1, 0]),
            "missing_numeric": int(np.isnan(att1_num).sum()),
            "negative_numeric": int((att1_num < 0).sum()),
            **cell_types(attach / "附件1.xlsx", 0),
        },
        "attachment2_load": stat(load),
        "attachment2_pv": stat(pv),
        "attachment4_price": stat(price),
        "cell_type_checks": {
            "attachment2_load": cell_types(attach / "附件2.xlsx", "小区负载"),
            "attachment2_pv": cell_types(attach / "附件2.xlsx", "光伏发电实际功率"),
            "attachment3": cell_types(attach / "附件3.xlsx", 0),
            "attachment4_price": cell_types(attach / "附件4.xlsx", 0),
        },
        "source_time_labels": {"count": len(source_labels), "first": source_labels[0], "last": source_labels[-1], "unique": len(set(source_labels))},
        "attachment3": {
            "rows": int(len(raw3)),
            "date_min": str(raw3["issue_date"].min().date()),
            "date_max": str(raw3["issue_date"].max().date()),
            "issue_hours": sorted(int(x) for x in raw3["issue_hour"].dropna().unique()),
            "dates_not_4_rows": int((issue_counts["count"] != 4).sum()),
            "dates_not_4_unique_issues": int((issue_counts["nunique"] != 4).sum()),
            "lead_columns": len(lead_cols),
            "forecast_missing": int(np.isnan(lead_values).sum()),
            "forecast_negative": int((lead_values < 0).sum()),
            "forecast_min_kw": float(np.nanmin(lead_values)),
            "forecast_max_kw": float(np.nanmax(lead_values)),
            "blank_date_cells_expected_ffill": int(raw3["_date_was_blank"].sum()),
            "target_timestamp_min": str(issue_long["target_timestamp"].min()),
            "target_timestamp_max": str(issue_long["target_timestamp"].max()),
            "targets_in_2026": int((issue_long["target_timestamp"] >= pd.Timestamp("2026-01-01")).sum()),
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="C题数据审计、因果预测与场景生成")
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = args.root.resolve()
    attach = root / "C题" / "附件"
    labels = read_output_labels(attach / "附件5" / "result2.xlsx")

    load, source_labels, _ = wide_sheet(attach / "附件2.xlsx", "小区负载")
    pv, source_labels_pv, _ = wide_sheet(attach / "附件2.xlsx", "光伏发电实际功率")
    dynamic_price, source_labels_price, _ = wide_sheet(attach / "附件4.xlsx", 0)
    if source_labels != source_labels_pv or source_labels != source_labels_price:
        raise ValueError("附件2/4时间表头不一致")
    if load.shape != (365, 144) or pv.shape != (365, 144) or dynamic_price.shape != (365, 144):
        raise ValueError("附件2/4应为365×144")
    if not (load.index.equals(pv.index) and load.index.equals(dynamic_price.index)):
        raise ValueError("附件2/4日期不一致")

    att1 = pd.read_excel(attach / "附件1.xlsx")
    static_price = pd.to_numeric(att1.iloc[:, 1], errors="coerce").to_numpy(float)
    if len(static_price) != 144:
        raise ValueError("附件1不是144个区间")

    raw3, issue_long = attachment3_long(attach / "附件3.xlsx")
    issue_long = issue_long.merge(pd.DataFrame({"interval_index": range(1, 145), "output_interval_label": labels}), on="interval_index", how="left")
    issue_long.to_csv(root / "scenario_data" / "official_forecasts_long.csv.gz", index=False, compression="gzip", encoding="utf-8", float_format="%.6f") if (root / "scenario_data").exists() else None

    idx = pd.MultiIndex.from_product([load.index, range(1, 145)], names=["date", "interval_index"])
    processed = idx.to_frame(index=False)
    processed["source_time_label"] = np.tile(source_labels, len(load))
    processed["output_interval_label"] = np.tile(labels, len(load))
    processed["interval_end_timestamp"] = processed["date"] + pd.to_timedelta(processed["interval_index"] * 10, unit="m")
    processed["load_actual_kw"] = load.to_numpy(float).reshape(-1)
    processed["pv_actual_kw"] = pv.to_numpy(float).reshape(-1)
    processed["price_actual"] = dynamic_price.to_numpy(float).reshape(-1)
    processed["price_static"] = np.tile(static_price, len(load))

    off0 = issue_long[(issue_long["issue_hour"] == 0) & (issue_long["issue_date"] == issue_long["date"])][["date", "interval_index", "forecast_issue_time", "forecast_lead_h", "pv_forecast_kw"]]
    off0 = off0.drop_duplicates(["date", "interval_index"])
    processed = processed.merge(off0, on=["date", "interval_index"], how="left")
    processed = processed[["date", "interval_index", "source_time_label", "output_interval_label", "interval_end_timestamp", "load_actual_kw", "pv_actual_kw", "price_actual", "price_static", "forecast_issue_time", "forecast_lead_h", "pv_forecast_kw"]]
    processed.to_csv(root / "processed_timeseries.csv", index=False, encoding="utf-8-sig", float_format="%.6f")

    forecasts, metrics = make_forecasts(processed, load.index, load, pv)
    metrics = pd.concat([metrics, official_issue_metrics(issue_long, processed)], ignore_index=True)
    forecasts.to_csv(root / "forecast_results.csv", index=False, encoding="utf-8-sig", float_format="%.6f")
    metrics.to_csv(root / "forecast_metrics.csv", index=False, encoding="utf-8-sig", float_format="%.6f")

    (root / "scenario_data").mkdir(exist_ok=True)
    issue_long.to_csv(root / "scenario_data" / "official_forecasts_long.csv.gz", index=False, compression="gzip", encoding="utf-8", float_format="%.6f")
    write_scenarios(root, forecasts, processed)
    make_provisional_solution(root, processed, forecasts, issue_long, labels)

    audit = build_audit(attach, load, pv, dynamic_price, raw3, issue_long, source_labels)
    audit_markdown(root, audit)
    data_dictionary(root)
    scenario_description(root)
    print(json.dumps({"status": "ok", "processed_rows": len(processed), "forecast_rows": len(forecasts), "metric_rows": len(metrics), "seed": SEED}, ensure_ascii=False))


if __name__ == "__main__":
    main()
