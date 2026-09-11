from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from export_official_templates import FILES, emergency_events, read_solution


EXPECTED_SHEETS = {
    "result1": ["计划购电量", "充放电量"],
    "result2": ["计划购电量", "充放电量", "紧急购电量"],
    "result3": ["计划购电量", "调整购电量", "充放电量", "紧急购电量"],
    "result4-2": ["计划购电量", "充放电量", "紧急购电量"],
    "result4-3": ["计划购电量", "调整购电量", "充放电量", "紧急购电量"],
}


def cached_number(ws, row: int, col: int) -> float | None:
    value = ws.cell(row, col).value
    return float(value) if isinstance(value, (int, float)) else None


def validate_file(workbook_root: Path, data_root: Path, name: str, solution: pd.DataFrame) -> tuple[dict[str, object], list[str]]:
    path = workbook_root / f"{name}.xlsx"
    wb_f = load_workbook(path, data_only=False, read_only=False)
    wb_v = load_workbook(path, data_only=True, read_only=False)
    g = solution[solution["strategy_name"].astype(str) == name].sort_values(["date", "interval_index"])
    issues: list[str] = []
    formula_errors = []
    for ws in wb_f.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    formula_errors.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
    if formula_errors:
        issues.append(f"公式错误单元格: {formula_errors[:10]}")

    details: dict[str, object] = {"file": name + ".xlsx", "sheets": wb_f.sheetnames, "formula_errors": len(formula_errors)}
    if wb_f.sheetnames != EXPECTED_SHEETS[name]:
        issues.append(f"工作表名称或顺序错误: {wb_f.sheetnames}")
    ellipsis = sum(
        str(cell.value).strip() in {"⁝", "…", "..."}
        for ws in wb_f.worksheets for row in ws.iter_rows() for cell in row if cell.value is not None
    )
    details["ellipsis_cells"] = ellipsis
    if ellipsis:
        issues.append(f"仍有{ellipsis}个省略号样例单元格")
    price_data = pd.read_csv(data_root / "processed_timeseries.csv", encoding="utf-8-sig", usecols=["date", "interval_index", "price_actual", "price_static"])
    price_data["date"] = pd.to_datetime(price_data["date"]).dt.normalize()
    cg = g.drop(columns=[c for c in ["price_actual", "price_static"] if c in g.columns]).merge(
        price_data, on=["date", "interval_index"], how="left"
    )
    price = cg["price_actual"].to_numpy(float) if name.startswith("result4-") else cg["price_static"].to_numpy(float)
    calc_plan = cg["purchase_plan_kwh"].to_numpy(float) * price
    if name in {"result3", "result4-3"}:
        diff = cg["purchase_adjust_kwh"].to_numpy(float) - cg["purchase_plan_kwh"].to_numpy(float)
        calc_adjust = np.where(diff >= 0, 1.5 * price * diff, 0.5 * price * (-diff))
    else:
        calc_adjust = np.zeros(len(cg))
    calc_emergency = 5.0 * price * cg["purchase_emergency_kwh"].to_numpy(float)
    cost_diff = max(
        float(np.max(np.abs(calc_plan - cg["cost_plan"].to_numpy(float)))),
        float(np.max(np.abs(calc_adjust - cg["cost_adjust"].to_numpy(float)))),
        float(np.max(np.abs(calc_emergency - cg["cost_emergency"].to_numpy(float)))),
        float(np.max(np.abs(calc_plan + calc_adjust + calc_emergency - cg["cost_total"].to_numpy(float)))),
    )
    details["independent_cost_max_diff"] = cost_diff
    if cost_diff > 1e-4:
        issues.append(f"按区间购电量和电价独立复算费用最大差{cost_diff}")
    if name == "result1":
        ws = wb_v["计划购电量"]
        if ws.cell(2, 1).value != "0:10-0:20" or ws.cell(145, 1).value != "0:00+1-0:10+1":
            issues.append("result1首末区间标签错误")
        actual = np.array([ws.cell(r, 2).value for r in range(2, 146)], dtype=float)
        expected = g["purchase_plan_kwh"].to_numpy(float)
        if not np.allclose(actual, expected, atol=1e-5):
            issues.append("result1计划购电量与solution_long不一致")
        details["intervals"] = len(actual)
        details["purchase_total_kwh"] = float(actual.sum())
        details["cost_total_yuan"] = float(g["cost_total"].sum())
        return details, issues

    dates = pd.date_range("2025-02-01", "2025-12-31", freq="D")
    ws = wb_v["计划购电量"]
    details["plan_shape"] = [ws.max_row, ws.max_column]
    if (ws.max_row, ws.max_column) != (335, 147):
        issues.append("计划购电量工作表不是335×147")
    if ws.cell(1, 2).value != "0:10-0:20" or ws.cell(1, 145).value != "0:00-0:10+1":
        issues.append("年度宽表首末区间标签错误")
    if pd.Timestamp(ws.cell(2, 1).value) != pd.Timestamp("2025-02-01") or pd.Timestamp(ws.cell(335, 1).value) != pd.Timestamp("2025-12-31"):
        issues.append("年度宽表首末日期错误")
    max_qty_diff = 0.0
    max_cost_diff = 0.0
    for r, d in enumerate(dates, 2):
        x = g[g["date"] == d]
        values = np.array([ws.cell(r, c).value for c in range(2, 146)], dtype=float)
        max_qty_diff = max(max_qty_diff, float(np.max(np.abs(values - x["purchase_plan_kwh"].to_numpy(float)))))
        total = cached_number(ws, r, 146)
        if total is None:
            issues.append(f"EP{r}公式无缓存值；需用Excel/LibreOffice重算")
            break
        max_qty_diff = max(max_qty_diff, abs(total - values.sum()))
        expected_cost = x["cost_total"].sum() if name in {"result2", "result4-2"} else x["cost_plan"].sum()
        max_cost_diff = max(max_cost_diff, abs(float(ws.cell(r, 147).value) - expected_cost))
    if max_qty_diff > 1e-4:
        issues.append(f"计划购电量/日合计最大差 {max_qty_diff}")
    if max_cost_diff > 1e-4:
        issues.append(f"计划购电费最大差 {max_cost_diff}")

    if "调整购电量" in wb_v.sheetnames:
        aw = wb_v["调整购电量"]
        adjust_diff = 0.0
        adjust_cost_diff = 0.0
        for r, d in enumerate(dates, 2):
            x = g[g["date"] == d]
            values = np.array([aw.cell(r, c).value for c in range(2, 146)], dtype=float)
            adjust_diff = max(adjust_diff, float(np.max(np.abs(values - x["purchase_adjust_kwh"].to_numpy(float)))))
            total = cached_number(aw, r, 146)
            if total is not None:
                adjust_diff = max(adjust_diff, abs(total - values.sum()))
            adjust_cost_diff = max(adjust_cost_diff, abs(float(aw.cell(r, 147).value) - x["cost_total"].sum()))
        if adjust_diff > 1e-4:
            issues.append(f"调整购电量/日合计最大差 {adjust_diff}")
        if adjust_cost_diff > 1e-4:
            issues.append(f"调整后全天费用最大差 {adjust_cost_diff}")
        details["max_adjust_diff"] = adjust_diff
        details["max_adjust_cost_diff"] = adjust_cost_diff

    cs = wb_v["充放电量"]
    if cs.max_row != 1 + 334 * 6:
        issues.append(f"充放电量工作表行数应为2005，实际{cs.max_row}")
    charge_dates = sum(cs.cell(r, 1).value is not None for r in range(2, cs.max_row + 1))
    details["charge_date_rows"] = charge_dates
    if charge_dates != 334:
        issues.append(f"充放电量日期行应为334，实际{charge_dates}")
    soc_continuity = 0.0
    by_day = [x.sort_values("interval_index") for _, x in g.groupby("date")]
    for a, b in zip(by_day[:-1], by_day[1:]):
        soc_continuity = max(soc_continuity, abs(float(a.iloc[-1]["soc_end_kwh"]) - float(b.iloc[0]["soc_start_kwh"])))
    if soc_continuity > 1e-4:
        issues.append(f"solution_long跨日SOC不连续，最大差{soc_continuity}")

    events = emergency_events(g)
    sheet_total = 0.0
    ew = wb_v["紧急购电量"]
    blank_rows = 0
    partial_rows = 0
    for r in range(2, ew.max_row + 1):
        vals = [ew.cell(r, c).value for c in range(1, 4)]
        blank_rows += int(all(v is None for v in vals))
        partial_rows += int((vals[1] is None) != (vals[2] is None))
        value = ew.cell(r, 3).value
        if isinstance(value, (int, float)):
            sheet_total += float(value)
    details["emergency_blank_rows"] = blank_rows
    details["emergency_partial_rows"] = partial_rows
    if blank_rows or partial_rows:
        issues.append(f"紧急购电表异常空行{blank_rows}、时间/数量不成对行{partial_rows}")
    expected_emergency = float(g["purchase_emergency_kwh"].sum())
    if abs(sheet_total - expected_emergency) > 1e-4:
        issues.append(f"紧急购电合并后总量不守恒: sheet={sheet_total}, solution={expected_emergency}")
    details.update({
        "max_plan_diff": max_qty_diff,
        "max_plan_cost_diff": max_cost_diff,
        "charge_rows": cs.max_row,
        "soc_crossday_max_diff": soc_continuity,
        "emergency_event_rows": len(events),
        "emergency_total_kwh": sheet_total,
    })
    return details, issues


def write_report(root: Path, solution_path: Path, reports: list[dict[str, object]], issues: dict[str, list[str]], provisional: bool) -> None:
    all_ok = not any(issues.values())
    status = "结构与数值校验通过，但仍为临时基线，待优化结果替换" if provisional and all_ok else ("全部校验通过" if all_ok else "存在待修复项")
    lines = [
        "# Excel 验证报告", "", f"结论：**{status}**。", "", f"输入：`{solution_path.name}`", "",
        "## 校验范围", "",
        "- 工作表名称与顺序保持官方模板不变。",
        "- 2025-02-01至2025-12-31共334天，每天144个区间。",
        "- 第t个值按官方模板第t列写入；未整体平移时间。",
        "- 充放电量展开为每天6个四小时段；省略号样例已删除。",
        "- 连续正紧急购电区间已合并，同日多段分行，合并前后总量守恒。",
        "- 全天购电量由Excel公式求和；费用与solution_long区间费用独立对比。",
        "- 文件保存后重新打开，并扫描公式错误字符串。", "",
        "## 文件结果", "",
        "| 文件 | 状态 | 关键统计 |", "|---|---|---|",
    ]
    for report in reports:
        name = report["file"]
        issue = issues.get(name, [])
        lines.append(f"| {name} | {'通过' if not issue else '待处理'} | `{json.dumps(report, ensure_ascii=False)}` |")
    if any(issues.values()):
        lines.extend(["", "## 待处理项", ""])
        for name, vals in issues.items():
            for value in vals:
                lines.append(f"- {name}: {value}")
    if provisional:
        lines.extend([
            "", "## 临时基线警告", "",
            "当前五个工作簿由`scenario_data/provisional_solution_long.csv`生成，用于验证完整导出链路。该基线严格使用因果预测，但采用零储能动作，`solver_status`含`PROVISIONAL`；不能作为最终竞赛最优结果。优化负责人提交最终`solution_long.csv`后，重新运行导出器和验证器。",
        ])
    (root / "excel_validation_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="保存后重开并校验五个官方结果文件")
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--workbook-root", type=Path, default=None)
    parser.add_argument("--data-root", type=Path, default=None)
    parser.add_argument("--solution", type=Path, default=None)
    args = parser.parse_args()
    root = args.root.resolve()
    workbook_root = args.workbook_root.resolve() if args.workbook_root else root
    data_root = args.data_root.resolve() if args.data_root else root
    solution_path = args.solution.resolve() if args.solution else data_root / "solution_long.csv"
    solution = read_solution(solution_path)
    reports = []
    issues: dict[str, list[str]] = {}
    for name in FILES:
        report, problem = validate_file(workbook_root, data_root, name, solution)
        reports.append(report)
        issues[report["file"]] = problem
    provisional = solution["solver_status"].astype(str).str.contains("PROVISIONAL", case=False).any()
    write_report(workbook_root, solution_path, reports, issues, provisional)
    print(json.dumps({"status": "ok" if not any(issues.values()) else "issues", "provisional": bool(provisional), "issues": issues}, ensure_ascii=False))


if __name__ == "__main__":
    main()
