from __future__ import annotations

import argparse
import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment


FILES = ["result1", "result2", "result3", "result4-2", "result4-3"]
REQUIRED = [
    "date", "interval_index", "output_interval_label", "purchase_plan_kwh",
    "purchase_adjust_kwh", "purchase_final_kwh", "purchase_emergency_kwh",
    "charge_kwh", "discharge_kwh", "soc_start_kwh", "soc_end_kwh",
    "curtailment_kwh", "cost_plan", "cost_adjust", "cost_emergency",
    "cost_total", "strategy_name", "solver_status",
]
SEGMENTS = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]


def copy_cell_style(src, dst) -> None:
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)


def capture_row(ws, row: int, cols: int) -> list[dict[str, object]]:
    out = []
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        out.append({"style": copy(cell._style), "number_format": cell.number_format, "alignment": copy(cell.alignment), "protection": copy(cell.protection)})
    return out


def apply_row(ws, row: int, styles: list[dict[str, object]]) -> None:
    for c, style in enumerate(styles, 1):
        cell = ws.cell(row, c)
        cell._style = copy(style["style"])
        cell.number_format = style["number_format"]
        cell.alignment = copy(style["alignment"])
        cell.protection = copy(style["protection"])


def read_solution(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, encoding="utf-8-sig")
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        raise ValueError(f"solution_long缺少字段: {missing}")
    df["date"] = pd.to_datetime(df["date"], errors="coerce").dt.normalize()
    df["interval_index"] = pd.to_numeric(df["interval_index"], errors="coerce")
    num_cols = [c for c in REQUIRED if c.endswith("_kwh") or c.startswith("cost_")]
    for c in num_cols:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    if df[["date", "interval_index"] + num_cols].isna().any().any():
        bad = df[["date", "interval_index"] + num_cols].isna().sum()
        raise ValueError(f"solution_long存在不可解析值: {bad[bad > 0].to_dict()}")
    dup = df.duplicated(["strategy_name", "date", "interval_index"]).sum()
    if dup:
        raise ValueError(f"solution_long重复主键: {dup}")
    if not df["interval_index"].between(1, 144).all():
        raise ValueError("interval_index必须为1..144")
    for c in ["purchase_plan_kwh", "purchase_adjust_kwh", "purchase_final_kwh", "purchase_emergency_kwh", "charge_kwh", "discharge_kwh"]:
        if (df[c] < -1e-7).any():
            raise ValueError(f"{c}出现负值")
    if (df["charge_kwh"] > 833.3333 + 1e-4).any() or (df["discharge_kwh"] > 833.3333 + 1e-4).any():
        raise ValueError("充/放电量超过每10分钟833.3333 kWh")
    if ((df[["soc_start_kwh", "soc_end_kwh"]] < 1200 - 1e-4).any().any()
            or (df[["soc_start_kwh", "soc_end_kwh"]] > 10800 + 1e-4).any().any()):
        raise ValueError("SOC越界")
    if ((df["charge_kwh"] > 1e-7) & (df["discharge_kwh"] > 1e-7)).any():
        raise ValueError("同一区间同时充电和放电")
    soc_expected = df["soc_start_kwh"] + 0.9 * df["charge_kwh"] - df["discharge_kwh"] / 0.9
    if (soc_expected - df["soc_end_kwh"]).abs().max() > 1e-3:
        raise ValueError("SOC状态方程不成立")
    if (df["cost_plan"] + df["cost_adjust"] + df["cost_emergency"] - df["cost_total"]).abs().max() > 1e-4:
        raise ValueError("cost_total不等于三个费用分量之和")
    for strategy, g in df.groupby("strategy_name"):
        g = g.sort_values(["date", "interval_index"])
        if len(g) > 1:
            same_run = (g["soc_start_kwh"].iloc[1:].to_numpy() - g["soc_end_kwh"].iloc[:-1].to_numpy())
            if np.abs(same_run).max() > 1e-3:
                raise ValueError(f"{strategy}区间SOC不连续")
        if strategy in {"result2", "result4-2"}:
            expected_final = g["purchase_plan_kwh"] + g["purchase_emergency_kwh"]
        elif strategy in {"result3", "result4-3"}:
            expected_final = g["purchase_adjust_kwh"] + g["purchase_emergency_kwh"]
        else:
            expected_final = g["purchase_plan_kwh"]
        if (g["purchase_final_kwh"] - expected_final).abs().max() > 1e-4:
            raise ValueError(f"{strategy}的purchase_final_kwh与采用的接口恒等式不符")
    return df.sort_values(["strategy_name", "date", "interval_index"]).reset_index(drop=True)


def template_labels(wb) -> list[str]:
    ws = wb["计划购电量"]
    return [str(ws.cell(1, c).value) for c in range(2, 146)]


def check_group(g: pd.DataFrame, labels: list[str], annual: bool) -> None:
    if annual:
        dates = pd.date_range("2025-02-01", "2025-12-31", freq="D")
        expected = len(dates) * 144
        if len(g) != expected or sorted(g["date"].unique()) != sorted(dates.to_numpy()):
            raise ValueError(f"年度策略必须覆盖334天×144={expected}行")
        counts = g.groupby("date")["interval_index"].nunique()
        if not (counts == 144).all():
            raise ValueError("存在非144区间日")
    else:
        if len(g) != 144 or g["interval_index"].nunique() != 144:
            raise ValueError("result1必须正好144行")
    if annual:
        expected_labels = np.tile(labels, g["date"].nunique())
        actual_labels = g.sort_values(["date", "interval_index"])["output_interval_label"].astype(str).to_numpy()
        if len(expected_labels) != len(actual_labels) or not np.array_equal(expected_labels, actual_labels):
            raise ValueError("output_interval_label与官方模板序号不一致")


def fill_wide(ws, g: pd.DataFrame, quantity: str, cost: str, labels: list[str]) -> None:
    dates = pd.date_range("2025-02-01", "2025-12-31", freq="D")
    by_date = {d: x.sort_values("interval_index") for d, x in g.groupby("date")}
    for r, d in enumerate(dates, 2):
        x = by_date[d]
        ws.cell(r, 1).value = d.to_pydatetime()
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
        for idx, value in enumerate(x[quantity].to_numpy(float), 2):
            ws.cell(r, idx).value = float(value)
        ws.cell(r, 146).value = f"=SUM(B{r}:EO{r})"
        ws.cell(r, 147).value = float(x[cost].sum())
        ws.cell(r, 147).comment = Comment(f"Source: solution_long.csv, sum({cost}) for {d.date()}", "Codex")
    for c, label in enumerate(labels, 2):
        if str(ws.cell(1, c).value) != label:
            raise ValueError(f"模板表头被改变: {ws.title}!{ws.cell(1,c).coordinate}")


def fill_q1(wb, g: pd.DataFrame) -> None:
    g = g.sort_values("interval_index")
    ws = wb["计划购电量"]
    for r, value in enumerate(g["purchase_plan_kwh"].to_numpy(float), 2):
        ws.cell(r, 2).value = float(value)
    cs = wb["充放电量"]
    for block in range(6):
        x = g[g["interval_index"].between(block * 24 + 1, (block + 1) * 24)]
        cs.cell(block + 2, 2).value = float(x["charge_kwh"].sum())
        cs.cell(block + 2, 3).value = float(x["discharge_kwh"].sum())
    cs.cell(2, 5).value = float(g.iloc[0]["soc_start_kwh"])
    cs.cell(3, 5).value = float(g.iloc[-1]["soc_end_kwh"])


def rebuild_charge_sheet(ws, g: pd.DataFrame) -> None:
    styles = [capture_row(ws, r, 6) for r in range(2, 8)]
    heights = [ws.row_dimensions[r].height for r in range(2, 8)]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    dates = pd.date_range("2025-02-01", "2025-12-31", freq="D")
    by_date = {d: x.sort_values("interval_index") for d, x in g.groupby("date")}
    row = 2
    for d in dates:
        x = by_date[d]
        for block, segment in enumerate(SEGMENTS):
            apply_row(ws, row, styles[block])
            ws.row_dimensions[row].height = heights[block]
            subset = x[x["interval_index"].between(block * 24 + 1, (block + 1) * 24)]
            ws.cell(row, 1).value = d.to_pydatetime() if block == 0 else None
            if block == 0:
                ws.cell(row, 1).number_format = "yyyy-mm-dd"
            ws.cell(row, 2).value = segment
            ws.cell(row, 3).value = float(subset["charge_kwh"].sum())
            ws.cell(row, 4).value = float(subset["discharge_kwh"].sum())
            ws.cell(row, 5).value = "0:00" if block == 0 else ("24:00" if block == 1 else None)
            ws.cell(row, 6).value = float(x.iloc[0]["soc_start_kwh"]) if block == 0 else (float(x.iloc[-1]["soc_end_kwh"]) if block == 1 else None)
            row += 1


def emergency_events(g: pd.DataFrame, tol: float = 1e-9) -> list[dict[str, object]]:
    events = []
    for d, x in g.groupby("date"):
        x = x.sort_values("interval_index")
        positive = x[x["purchase_emergency_kwh"] > tol]
        if positive.empty:
            continue
        block = []
        blocks = []
        prev = None
        for _, row in positive.iterrows():
            idx = int(row["interval_index"])
            if prev is None or idx == prev + 1:
                block.append(row)
            else:
                blocks.append(block)
                block = [row]
            prev = idx
        if block:
            blocks.append(block)
        for rows in blocks:
            first, last = rows[0], rows[-1]
            start = str(first["output_interval_label"]).split("-", 1)[0]
            end = str(last["output_interval_label"]).split("-", 1)[1]
            events.append({"date": d, "period": f"{start}-{end}", "amount": float(sum(float(r["purchase_emergency_kwh"]) for r in rows))})
    return events


def rebuild_emergency_sheet(ws, g: pd.DataFrame) -> None:
    style = capture_row(ws, 2, 3)
    height = ws.row_dimensions[2].height
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    events = emergency_events(g)
    row = 2
    last_date = None
    for event in events:
        apply_row(ws, row, style)
        ws.row_dimensions[row].height = height
        if event["date"] != last_date:
            ws.cell(row, 1).value = pd.Timestamp(event["date"]).to_pydatetime()
            ws.cell(row, 1).number_format = "yyyy-mm-dd"
        ws.cell(row, 2).value = event["period"]
        ws.cell(row, 3).value = event["amount"]
        last_date = event["date"]
        row += 1


def export_one(name: str, template_dir: Path, out_dir: Path, solution: pd.DataFrame) -> dict[str, object]:
    src = template_dir / f"{name}.xlsx"
    dst = out_dir / f"{name}.xlsx"
    shutil.copy2(src, dst)
    wb = load_workbook(dst)
    original_sheets = list(wb.sheetnames)
    labels = template_labels(wb)
    g = solution[solution["strategy_name"].astype(str) == name].copy()
    check_group(g, labels, annual=name != "result1")

    if name == "result1":
        fill_q1(wb, g)
    else:
        plan_cost = "cost_total" if name in {"result2", "result4-2"} else "cost_plan"
        fill_wide(wb["计划购电量"], g, "purchase_plan_kwh", plan_cost, labels)
        if "调整购电量" in wb.sheetnames:
            fill_wide(wb["调整购电量"], g, "purchase_adjust_kwh", "cost_total", labels)
        rebuild_charge_sheet(wb["充放电量"], g)
        rebuild_emergency_sheet(wb["紧急购电量"], g)

    if wb.sheetnames != original_sheets:
        raise ValueError(f"{name}工作表顺序发生变化")
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(dst)

    reopened = load_workbook(dst, data_only=False, read_only=False)
    if reopened.sheetnames != original_sheets:
        raise ValueError(f"{name}保存后工作表顺序不一致")
    ellipsis = 0
    for ws in reopened.worksheets:
        for row in ws.iter_rows():
            ellipsis += sum(str(c.value).strip() in {"⁝", "…", "..."} for c in row if c.value is not None)
    if ellipsis:
        raise ValueError(f"{name}仍含省略号示例")
    return {"file": name + ".xlsx", "sheets": original_sheets, "rows": {ws.title: ws.max_row for ws in reopened.worksheets}, "ellipsis_cells": ellipsis}


def main() -> None:
    parser = argparse.ArgumentParser(description="从官方模板副本导出五个结果工作簿")
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--solution", type=Path, default=None)
    args = parser.parse_args()
    root = args.root.resolve()
    solution_path = args.solution.resolve() if args.solution else root / "solution_long.csv"
    solution = read_solution(solution_path)
    missing_strategies = [x for x in FILES if x not in set(solution["strategy_name"].astype(str))]
    if missing_strategies:
        raise ValueError(f"缺少策略: {missing_strategies}")
    template_dir = root / "source_inputs" / "附件5"
    output_dir = args.output_dir.resolve() if args.output_dir else root / "official_results"
    output_dir.mkdir(parents=True, exist_ok=True)
    reports = [export_one(name, template_dir, output_dir, solution) for name in FILES]
    print(json.dumps({"status": "ok", "solution": str(solution_path), "output_dir": str(output_dir), "files": reports}, ensure_ascii=False))


if __name__ == "__main__":
    main()
